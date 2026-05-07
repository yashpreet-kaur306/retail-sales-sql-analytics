import os
import sys
import sqlite3
import subprocess
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (PatternFill, Font, Alignment,
                              Border, Side)
from openpyxl.utils import get_column_letter
from datetime import datetime

# ── PATHS ──
BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
DB_PATH     = os.path.join(BASE_DIR, "retail_sales.db")
SQL_PATH    = os.path.join(BASE_DIR, "sql", "analysis_queries.sql")
DATA_SCRIPT = os.path.join(BASE_DIR, "data", "generate_data.py")
REPORTS_DIR = os.path.join(BASE_DIR, "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


if not os.path.exists(DB_PATH):
    print("📦  Database not found. Generating data...")
    subprocess.run([sys.executable, DATA_SCRIPT], check=True)
else:
    print("✅  Database found. Skipping data generation.")


conn = sqlite3.connect(DB_PATH)

QUERIES = {
    "Top Products": """
        SELECT p.product_id, p.product_name, p.category,
               COUNT(o.order_id) AS total_orders,
               SUM(o.quantity)   AS total_units_sold,
               ROUND(SUM(o.total_amount), 2) AS total_revenue,
               ROUND(AVG(o.total_amount), 2) AS avg_order_value
        FROM orders o JOIN products p ON o.product_id = p.product_id
        GROUP BY p.product_id, p.product_name, p.category
        ORDER BY total_revenue DESC
    """,
    "Monthly Revenue": """
        SELECT strftime('%Y', order_date) AS year,
               strftime('%m', order_date) AS month,
               COUNT(order_id) AS total_orders,
               SUM(quantity)   AS total_units,
               ROUND(SUM(total_amount), 2) AS monthly_revenue
        FROM orders
        GROUP BY year, month ORDER BY year, month
    """,
    "Regional Breakdown": """
        SELECT region,
               COUNT(order_id) AS total_orders,
               SUM(quantity)   AS total_units,
               ROUND(SUM(total_amount), 2) AS regional_revenue,
               ROUND(SUM(total_amount)*100.0/(SELECT SUM(total_amount) FROM orders),2)
                   AS revenue_share_pct
        FROM orders
        GROUP BY region ORDER BY regional_revenue DESC
    """,
    "Customer Retention": """
        WITH customer_years AS (
            SELECT customer_id,
                   COUNT(DISTINCT strftime('%Y', order_date)) AS active_years,
                   COUNT(order_id) AS total_orders,
                   ROUND(SUM(total_amount), 2) AS lifetime_value
            FROM orders GROUP BY customer_id
        )
        SELECT cy.customer_id, c.customer_name, c.region,
               cy.active_years, cy.total_orders, cy.lifetime_value,
               CASE WHEN cy.active_years >= 3 THEN 'Loyal'
                    WHEN cy.active_years = 2  THEN 'Returning'
                    ELSE 'New' END AS customer_segment
        FROM customer_years cy
        JOIN customers c ON cy.customer_id = c.customer_id
        ORDER BY cy.lifetime_value DESC
    """,
    "Category Performance": """
        WITH category_stats AS (
            SELECT p.category,
                   COUNT(o.order_id) AS total_orders,
                   SUM(o.quantity)   AS total_units,
                   ROUND(SUM(o.total_amount),2) AS category_revenue
            FROM orders o JOIN products p ON o.product_id = p.product_id
            GROUP BY p.category
        )
        SELECT category, total_orders, total_units, category_revenue,
               ROUND(SUM(category_revenue) OVER (ORDER BY category_revenue DESC),2)
                   AS running_total,
               ROUND(category_revenue*100.0/SUM(category_revenue) OVER (),2)
                   AS revenue_share_pct
        FROM category_stats ORDER BY category_revenue DESC
    """,
    "Top Customers by Region": """
        WITH cr AS (
            SELECT o.customer_id, c.customer_name, o.region,
                   ROUND(SUM(o.total_amount),2) AS total_spent,
                   COUNT(o.order_id) AS total_orders
            FROM orders o JOIN customers c ON o.customer_id=c.customer_id
            GROUP BY o.customer_id, c.customer_name, o.region
        ),
        ranked AS (
            SELECT *, RANK() OVER (PARTITION BY region ORDER BY total_spent DESC)
                AS region_rank FROM cr
        )
        SELECT * FROM ranked WHERE region_rank <= 5
        ORDER BY region, region_rank
    """,
    "MoM Growth": """
        WITH monthly AS (
            SELECT strftime('%Y-%m', order_date) AS ym,
                   ROUND(SUM(total_amount),2) AS revenue
            FROM orders GROUP BY ym
        )
        SELECT ym, revenue,
               LAG(revenue) OVER (ORDER BY ym) AS prev_month_revenue,
               ROUND((revenue - LAG(revenue) OVER (ORDER BY ym))*100.0/
                     NULLIF(LAG(revenue) OVER (ORDER BY ym),0),2) AS mom_growth_pct
        FROM monthly ORDER BY ym
    """,
    "Discount Impact": """
        SELECT CASE WHEN discount_pct=0   THEN 'No Discount'
                    WHEN discount_pct<=5  THEN 'Low (1-5%)'
                    WHEN discount_pct<=10 THEN 'Medium (6-10%)'
                    ELSE 'High (>10%)' END AS discount_tier,
               COUNT(order_id) AS total_orders,
               ROUND(AVG(total_amount),2) AS avg_order_value,
               ROUND(SUM(total_amount),2) AS total_revenue
        FROM orders
        GROUP BY discount_tier ORDER BY avg_order_value DESC
    """,
}

results = {}
for name, query in QUERIES.items():
    print(f"  ⚙️  Running: {name}...")
    results[name] = pd.read_sql_query(query, conn)

conn.close()
print(f"\n✅  All {len(results)} queries executed successfully.\n")


timestamp   = datetime.now().strftime("%Y%m%d_%H%M%S")
output_path = os.path.join(REPORTS_DIR, f"Retail_MIS_Report_{timestamp}.xlsx")

with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    for sheet_name, df in results.items():
        short = sheet_name[:31]          # Excel sheet name limit
        df.to_excel(writer, sheet_name=short, index=False)


wb = load_workbook(output_path)

HEADER_FILL  = PatternFill("solid", fgColor="1F4E79")
ALT_FILL     = PatternFill("solid", fgColor="D6E4F0")
HEADER_FONT  = Font(bold=True, color="FFFFFF", size=11)
BODY_FONT    = Font(size=10)
BORDER_SIDE  = Side(style="thin", color="B0C4DE")
CELL_BORDER  = Border(left=BORDER_SIDE, right=BORDER_SIDE,
                      top=BORDER_SIDE,  bottom=BORDER_SIDE)

for ws in wb.worksheets:
    # header row
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = CELL_BORDER

    ws.row_dimensions[1].height = 28

    # data rows
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        fill = ALT_FILL if row_idx % 2 == 0 else PatternFill()
        for cell in row:
            cell.fill      = fill
            cell.font      = BODY_FONT
            cell.alignment = Alignment(horizontal="center")
            cell.border    = CELL_BORDER

    # auto-width columns
    for col in ws.columns:
        max_len = max(
            (len(str(c.value)) for c in col if c.value is not None),
            default=10
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    # freeze header row
    ws.freeze_panes = "A2"

wb.save(output_path)
print(f"📊  Excel MIS report saved → {output_path}")
print("\n🎉  Pipeline complete!")